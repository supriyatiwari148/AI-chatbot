"""
AI Voice Assistant - Flask backend.

Responsibilities:
  * Serve the web page (templates/index.html + static files)
  * Receive chat messages from the browser
  * Forward the conversation to a local Ollama model (text model, or a vision model when a photo is attached)
  * Read uploaded documents (PDF / DOCX / text) and return their text
  * Return the AI reply (or a friendly error) as JSON
"""

import base64
import binascii
import io
import os

import requests
from flask import Flask, jsonify, render_template, request

# ---------------------------------------------------------------------------
# Configuration (can be overridden with environment variables)
# ---------------------------------------------------------------------------
OLLAMA_CHAT_URL = os.environ.get("OLLAMA_URL", "http://localhost:11434/api/chat")
OLLAMA_MODEL = os.environ.get("OLLAMA_MODEL", "llama3.2:3b")
# Vision model: used automatically only when the user attaches a photo.
# llama3.2:3b cannot see images, so a second (multimodal) model is needed.
# Lighter alternative for weak laptops:  set OLLAMA_VISION_MODEL=moondream
OLLAMA_VISION_MODEL = os.environ.get("OLLAMA_VISION_MODEL", "llava")

# Cloud mode: if GROQ_API_KEY is set (e.g. on Render), the app uses Groq's free
# API instead of a local Ollama. Without a key it keeps using Ollama locally.
GROQ_API_KEY = os.environ.get("GROQ_API_KEY", "").strip()
GROQ_MODEL = os.environ.get("GROQ_MODEL", "llama-3.1-8b-instant")
# Groq vision model (check console.groq.com/docs/vision for the current name)
GROQ_VISION_MODEL = os.environ.get("GROQ_VISION_MODEL", "meta-llama/llama-4-scout-17b-16e-instruct")
GROQ_CHAT_URL = "https://api.groq.com/openai/v1/chat/completions"
USE_GROQ = bool(GROQ_API_KEY)
ACTIVE_MODEL = GROQ_MODEL if USE_GROQ else OLLAMA_MODEL
REQUEST_TIMEOUT_SECONDS = 180      # small models on a laptop CPU can be slow
MAX_MESSAGE_LENGTH = 2000          # characters allowed in one user message
MAX_HISTORY_MESSAGES = 20          # how many previous messages we send to the AI
MAX_FILE_CHARS = 12000             # characters of a document sent to the AI (small models have small memory)
MAX_FILES_PER_MESSAGE = 3
MAX_IMAGES_PER_MESSAGE = 4
MAX_IMAGE_BASE64_CHARS = 6 * 1024 * 1024   # ~4.5 MB image after base64
MAX_UPLOAD_BYTES = 10 * 1024 * 1024        # 10 MB per document upload
OLLAMA_NUM_CTX = 8192              # context window; Ollama's default (2-4k) would cut off documents

SYSTEM_PROMPT = (
    "You are AI Voice Assistant, a friendly and helpful assistant. "
    "Explain things clearly in simple language, like a good teacher. "
    "Your answers may be read aloud, so keep them concise and avoid large tables. "
    "Put any code inside Markdown code blocks. "
    "Use the earlier messages in the conversation to understand follow-up questions."
)

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 20 * 1024 * 1024  # reject request bodies over 20 MB (images/documents)


# ---------------------------------------------------------------------------
# Helper functions
# ---------------------------------------------------------------------------
def error_response(message, status_code):
    """Return a clean JSON error."""
    return jsonify({"error": message}), status_code


def clean_images(raw_images):
    """Validate a list of base64 images (the browser sends JPEGs without the data: prefix)."""
    if not isinstance(raw_images, list):
        return []
    cleaned = []
    for image in raw_images[:MAX_IMAGES_PER_MESSAGE]:
        if not isinstance(image, str):
            continue
        if image.startswith("data:") and "," in image:      # tolerate full data URLs
            image = image.split(",", 1)[1]
        image = image.strip()
        if not image or len(image) > MAX_IMAGE_BASE64_CHARS:
            continue
        try:
            base64.b64decode(image, validate=True)
        except (binascii.Error, ValueError):
            continue
        cleaned.append(image)
    return cleaned


def clean_history(raw_history):
    """Keep only valid {role, content} items and only the most recent ones.

    Earlier user messages may also carry "images" (the browser only re-sends the latest photo).
    """
    if not isinstance(raw_history, list):
        return []
    cleaned = []
    for item in raw_history:
        if not isinstance(item, dict):
            continue
        role = item.get("role")
        content = item.get("content")
        if role in ("user", "assistant") and isinstance(content, str) and content.strip():
            entry = {"role": role, "content": content.strip()[: MAX_FILE_CHARS * 2 + MAX_MESSAGE_LENGTH]}
            images = clean_images(item.get("images")) if role == "user" else []
            if images:
                entry["images"] = images
            cleaned.append(entry)
    return cleaned[-MAX_HISTORY_MESSAGES:]


def clean_files(raw_files):
    """Validate [{name, text}] documents whose text was extracted in the browser."""
    if not isinstance(raw_files, list):
        return []
    cleaned = []
    for item in raw_files[:MAX_FILES_PER_MESSAGE]:
        if not isinstance(item, dict):
            continue
        name = item.get("name")
        text = item.get("text")
        if isinstance(name, str) and isinstance(text, str) and text.strip():
            cleaned.append({"name": name.strip()[:120] or "file", "text": text.strip()[:MAX_FILE_CHARS]})
    return cleaned


def compose_user_text(message, files):
    """Put the document text in front of the user's question so the model can read it."""
    parts = []
    for item in files:
        parts.append(f"[Attached file: {item['name']}]\n```\n{item['text']}\n```")
    parts.append(message)
    return "\n\n".join(parts)


def extract_pdf_text(data):
    from pypdf import PdfReader   # imported here so the app still starts if pypdf is missing
    reader = PdfReader(io.BytesIO(data))
    pages = []
    for page in reader.pages:
        pages.append(page.extract_text() or "")
        if sum(len(p) for p in pages) > MAX_FILE_CHARS:
            break
    return "\n".join(pages)


def extract_docx_text(data):
    import docx                    # python-docx
    document = docx.Document(io.BytesIO(data))
    lines = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(lines)


def pick_model(messages):
    """Use the vision model only when at least one message contains an image."""
    has_images = any(m.get("images") for m in messages)
    if USE_GROQ:
        return (GROQ_VISION_MODEL if has_images else GROQ_MODEL), has_images
    return (OLLAMA_VISION_MODEL if has_images else OLLAMA_MODEL), has_images


def ask_ollama(messages, model):
    """Send the messages to Ollama and return the response.

    Ollama takes images as a plain list of base64 strings on the message.
    Raises requests exceptions, which the route handles.
    """
    payload = {
        "model": model,
        "messages": [{"role": "system", "content": SYSTEM_PROMPT}] + messages,
        "stream": False,
        "options": {"num_ctx": OLLAMA_NUM_CTX},
    }
    return requests.post(OLLAMA_CHAT_URL, json=payload, timeout=REQUEST_TIMEOUT_SECONDS)


def to_openai_messages(messages):
    """Convert our messages to the OpenAI/Groq format (images become data URLs)."""
    converted = []
    for m in messages:
        if m.get("images"):
            parts = [{"type": "text", "text": m["content"]}]
            for image in m["images"]:
                parts.append({"type": "image_url", "image_url": {"url": "data:image/jpeg;base64," + image}})
            converted.append({"role": m["role"], "content": parts})
        else:
            converted.append({"role": m["role"], "content": m["content"]})
    return converted


def ask_groq(messages, model):
    """Send the messages to Groq (OpenAI-compatible API) and return the response."""
    payload = {
        "model": model,
        "messages": [{"role": "system", "content": SYSTEM_PROMPT}] + to_openai_messages(messages),
        "stream": False,
    }
    headers = {"Authorization": f"Bearer {GROQ_API_KEY}"}
    return requests.post(GROQ_CHAT_URL, json=payload, headers=headers, timeout=REQUEST_TIMEOUT_SECONDS)


# ---------------------------------------------------------------------------
# Security headers
# ---------------------------------------------------------------------------
@app.after_request
def add_security_headers(response):
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["X-Frame-Options"] = "DENY"
    response.headers["Referrer-Policy"] = "no-referrer"
    response.headers["Content-Security-Policy"] = (
        "default-src 'self'; script-src 'self'; style-src 'self'; "
        "img-src 'self' data:; connect-src 'self'; object-src 'none'; base-uri 'self'"
    )
    return response


# ---------------------------------------------------------------------------
# Routes
# ---------------------------------------------------------------------------
@app.route("/")
def index():
    return render_template("index.html", model_name=ACTIVE_MODEL, cloud_mode=USE_GROQ)


@app.route("/api/health")
def health():
    """Tell the frontend whether the AI backend is reachable and the models are ready."""
    if USE_GROQ:
        return jsonify({"ollama": True, "model": GROQ_MODEL, "model_installed": True,
                        "vision_model": GROQ_VISION_MODEL, "vision_installed": True})
    try:
        tags_url = OLLAMA_CHAT_URL.replace("/api/chat", "/api/tags")
        response = requests.get(tags_url, timeout=5)
        installed = [m.get("name", "") for m in response.json().get("models", [])]

        def is_installed(wanted):
            return any(name == wanted or name.startswith(wanted + ":") for name in installed)

        return jsonify({
            "ollama": True,
            "model": OLLAMA_MODEL, "model_installed": is_installed(OLLAMA_MODEL),
            "vision_model": OLLAMA_VISION_MODEL, "vision_installed": is_installed(OLLAMA_VISION_MODEL),
        })
    except (requests.exceptions.RequestException, ValueError):
        return jsonify({"ollama": False, "model": OLLAMA_MODEL, "model_installed": False,
                        "vision_model": OLLAMA_VISION_MODEL, "vision_installed": False})


@app.route("/api/extract", methods=["POST"])
def extract():
    """Read a document (PDF, DOCX or plain text/code) and return its text to the browser."""
    upload = request.files.get("file")
    if upload is None or not upload.filename:
        return error_response("No file was uploaded.", 400)

    data = upload.read(MAX_UPLOAD_BYTES + 1)
    if len(data) > MAX_UPLOAD_BYTES:
        return error_response("That file is too large (maximum 10 MB).", 413)
    if not data:
        return error_response("That file is empty.", 400)

    name = upload.filename
    lower = name.lower()
    try:
        if lower.endswith(".pdf"):
            text = extract_pdf_text(data)
        elif lower.endswith(".docx"):
            text = extract_docx_text(data)
        else:
            if b"\x00" in data[:4096]:
                return error_response("This file type is not supported. Try PDF, DOCX, TXT, code or CSV files.", 415)
            text = data.decode("utf-8", errors="replace")
    except ImportError:
        return error_response("Server is missing a library for this file type. Run: pip install -r requirements.txt", 500)
    except Exception as error:   # corrupted / password protected files raise many different errors
        print("Extract error:", type(error).__name__, error)
        return error_response("Could not read that file. It may be corrupted or password-protected.", 422)

    text = text.strip()
    if not text:
        return error_response("No readable text found in that file (scanned PDFs are not supported).", 422)

    truncated = len(text) > MAX_FILE_CHARS
    return jsonify({"name": name, "text": text[:MAX_FILE_CHARS], "truncated": truncated})


@app.route("/api/chat", methods=["POST"])
def chat():
    data = request.get_json(silent=True)
    if not isinstance(data, dict):
        return error_response("Invalid request. Please send JSON.", 400)

    user_message = data.get("message")
    user_message = user_message.strip() if isinstance(user_message, str) else ""
    if len(user_message) > MAX_MESSAGE_LENGTH:
        return error_response(f"Message is too long (maximum {MAX_MESSAGE_LENGTH} characters).", 400)

    files = clean_files(data.get("files"))
    images = clean_images(data.get("images"))

    if not user_message:
        if images:
            user_message = "Describe this image in detail."
        elif files:
            user_message = "Summarize this file."
        else:
            return error_response("Please type a message before sending.", 400)

    new_message = {"role": "user", "content": compose_user_text(user_message, files)}
    if images:
        new_message["images"] = images

    # Previous messages + the new message = full conversation context
    messages = clean_history(data.get("history")) + [new_message]
    model, has_images = pick_model(messages)

    if USE_GROQ:
        try:
            ai_response = ask_groq(messages, model)
        except requests.exceptions.Timeout:
            return error_response("The AI took too long to answer. Please try again.", 504)
        except requests.exceptions.RequestException:
            return error_response("Could not reach the AI service. Please try again.", 502)

        if ai_response.status_code == 401:
            return error_response("The AI service rejected the API key. Check GROQ_API_KEY.", 502)
        if ai_response.status_code == 429:
            return error_response("Too many requests right now. Please wait a moment and try again.", 429)
        if ai_response.status_code != 200:
            print("Groq error:", ai_response.status_code, ai_response.text[:300])
            if has_images:
                return error_response("The vision model could not read the image. Check GROQ_VISION_MODEL.", 502)
            return error_response("The AI service returned an error. Please try again.", 502)

        try:
            reply = ai_response.json()["choices"][0]["message"]["content"].strip()
        except (ValueError, KeyError, IndexError, TypeError, AttributeError):
            return error_response("Received an invalid response from the AI service.", 502)
    else:
        try:
            ollama_response = ask_ollama(messages, model)
        except requests.exceptions.ConnectionError:
            return error_response(
                "Cannot reach Ollama. Make sure the Ollama app is running (open it from the Start menu).", 503
            )
        except requests.exceptions.Timeout:
            return error_response("The AI took too long to answer. Try a shorter question or a smaller model.", 504)
        except requests.exceptions.RequestException:
            return error_response("Something went wrong while contacting the AI model.", 502)

        if ollama_response.status_code == 404:
            return error_response(f"The model '{model}' is not installed. Run:  ollama pull {model}", 404)
        if ollama_response.status_code != 200:
            try:
                detail = ollama_response.json().get("error", "")
            except ValueError:
                detail = ollama_response.text[:200]
            print("Ollama error:", ollama_response.status_code, detail)
            return error_response(f"The AI model returned an error: {detail}", 502)

        try:
            reply = ollama_response.json()["message"]["content"].strip()
        except (ValueError, KeyError, TypeError, AttributeError):
            return error_response("Received an invalid response from the AI model.", 502)

    if not reply:
        return error_response("The AI returned an empty answer. Please try again.", 502)

    return jsonify({"reply": reply, "model": model})


# ---------------------------------------------------------------------------
# JSON error handlers (so the browser never receives raw Python error pages)
# ---------------------------------------------------------------------------
@app.errorhandler(413)
def too_large(_error):
    return error_response("That file or image is too large (maximum about 10 MB).", 413)


@app.errorhandler(405)
def method_not_allowed(_error):
    return error_response("Method not allowed.", 405)


@app.errorhandler(500)
def server_error(_error):
    return error_response("Internal server error. Please try again.", 500)


if __name__ == "__main__":
    debug_mode = os.environ.get("FLASK_DEBUG", "0") == "1"
    app.run(host="127.0.0.1", port=int(os.environ.get("PORT", 5000)), debug=debug_mode)
